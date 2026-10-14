#!/usr/bin/env python3
"""Buat presentasi .pptx yang merangkum project Gemastik Data Mining.

Jalankan dari root project:
    python tools/generate_presentation_pptx.py

Output: presentation.pptx
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN


OUT = Path("presentation.pptx")
FIGURES = Path("figures")


def _add_title_slide(prs: Presentation, title: str, subtitle: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[0])  # Title Slide
    slide.shapes.title.text = title
    slide.placeholders[1].text = subtitle


def _add_content_slide(
    prs: Presentation,
    title: str,
    bullets: list[str],
    image_path: Path | None = None,
) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[1])  # Title + Content
    slide.shapes.title.text = title
    tf = slide.placeholders[1].text_frame
    for i, b in enumerate(bullets):
        if i == 0:
            tf.text = b
        else:
            p = tf.add_paragraph()
            p.text = b
    if image_path and image_path.exists():
        slide.shapes.add_picture(str(image_path), Inches(5.5), Inches(1.8), width=Inches(4))


def build() -> Presentation:
    prs = Presentation()

    # 1 – Judul
    _add_title_slide(
        prs,
        "Klasifikasi Teks Bahasa Indonesia",
        "Gemastik – Data Mining\n(isi nama tim)",
    )

    # 2 – Latar Belakang & Rumusan Masalah
    _add_content_slide(prs, "Latar Belakang & Rumusan Masalah", [
        "Internet penuh dengan komentar toxic (kasar/negatif).",
        "Moderasi manual tidak efisien untuk volume data besar.",
        "Tujuan: Membangun sistem otomatis untuk mendeteksi teks toxic Bahasa Indonesia.",
        "Rumusan masalah: Bagaimana mengklasifikasikan teks ke label 0 (positif) atau 1 (toxic)?",
        "Metrik evaluasi: Macro F1-Score.",
    ])

    # 3 – Deskripsi Data
    _add_content_slide(prs, "Deskripsi Data", [
        "Data disediakan oleh panitia kompetisi Gemastik:",
        "• labeled.csv — 800 baris (teks + label) untuk training",
        "• unlabeled.csv — 2.400 baris (teks saja) untuk memperkaya kosakata",
        "• test.csv — 400 baris yang harus diprediksi",
        "Distribusi label cukup seimbang: 52.4% (label 0) vs 47.6% (label 1).",
    ])

    # 4 – EDA: Distribusi Kelas
    _add_content_slide(prs, "EDA — Distribusi Kelas", [
        "Dataset berlabel berjumlah 800 baris.",
        "Label 0 (positif): 419 baris (52.4%).",
        "Label 1 (toxic): 381 baris (47.6%).",
        "Distribusi relatif seimbang → tidak perlu teknik oversampling.",
    ], FIGURES / "01_class_distribution.png")

    # 5 – EDA: Panjang Teks
    _add_content_slide(prs, "EDA — Distribusi Panjang Teks", [
        "Rata-rata: ~25 kata per dokumen.",
        "Median: 19 kata (sebagian besar teks pendek).",
        "Teks pendek cocok untuk pendekatan TF-IDF + linear model.",
    ], FIGURES / "02_word_length_hist.png")

    # 6 – EDA: Top Bigrams
    _add_content_slide(prs, "EDA — Top Bigrams", [
        "Bigram = pasangan dua kata berurutan.",
        "Bigram paling sering menunjukkan topik & pola dominan.",
        "Membantu memahami karakteristik teks toxic vs positif.",
    ], FIGURES / "04_top_ngrams.png")

    # 7 – Preprocessing
    _add_content_slide(prs, "Preprocessing Teks", [
        "1. Lowercasing — huruf kecil semua.",
        "2. Hapus noise — URL, mention (@), hashtag (#), angka, tanda baca.",
        "3. Normalisasi slang — kamus 70+ kata (gak→tidak, gw→saya, bgt→sangat).",
        "4. Squash huruf berulang — baguuuus → bagus.",
        "Tujuan: Menyeragamkan teks agar model tidak bingung dengan variasi penulisan.",
    ])

    # 8 – Feature Engineering
    _add_content_slide(prs, "Feature Engineering — TF-IDF", [
        "TF-IDF mengubah teks menjadi representasi angka.",
        "Kata penting (sering di satu dokumen, jarang di lainnya) → nilai tinggi.",
        "Pengaturan: 50.000 fitur, unigram + bigram, sublinear_tf.",
        "Transductive: fit pada labeled + unlabeled + test (hanya frekuensi, bukan label).",
        "Hasilnya: kosakata lebih kaya, terutama untuk kata yang hanya ada di test set.",
    ])

    # 9 – Metodologi Pemodelan
    _add_content_slide(prs, "Metodologi Pemodelan", [
        "3 model diuji: Multinomial NB, Logistic Regression, Linear SVC.",
        "Evaluasi: 5-Fold Stratified Cross-Validation.",
        "Stratified = proporsi label dipertahankan di setiap fold.",
        "Metrik: Macro F1-Score (rata-rata F1 per kelas).",
        "Kenapa model linear? Dataset kecil (800) → deep learning rentan overfitting.",
    ])

    # 10 – Hasil
    _add_content_slide(prs, "Hasil Eksperimen", [
        "MultinomialNB          — F1 = 0.8486 ± 0.0178",
        "LogisticRegression    — F1 = 0.8674 ± 0.0218",
        "LinearSVC (terbaik!)  — F1 = 0.8734 ± 0.0146",
        "",
        "LinearSVC dipilih karena F1 tertinggi dan standar deviasi terendah (paling stabil).",
    ])

    # 11 – Analisis Hasil
    _add_content_slide(prs, "Analisis Hasil", [
        "Mengapa LinearSVC terbaik?",
        "• Optimal untuk data sparse berdimensi tinggi (TF-IDF).",
        "• Margin-based: mencari batas pemisah yang paling optimal.",
        "• Stabil di berbagai fold (std = 0.0146).",
        "",
        "Mengapa pipeline ini berhasil?",
        "• Normalisasi slang menyeragamkan kosakata.",
        "• Transductive TF-IDF memperkaya representasi fitur.",
    ])

    # 12 – Kesimpulan
    _add_content_slide(prs, "Kesimpulan", [
        "Pipeline TF-IDF + Linear SVC mencapai F1 ≈ 0.87.",
        "Pendekatan klasik NLP terbukti efektif untuk dataset kecil.",
        "Normalisasi slang Bahasa Indonesia sangat membantu performa.",
        "",
        "Potensi perbaikan:",
        "• Semi-supervised learning (manfaatkan unlabeled data).",
        "• Word embeddings (FastText/IndoBERT).",
        "• Ensemble methods.",
    ])

    # 13 – Terima Kasih
    _add_title_slide(
        prs,
        "Terima Kasih",
        "Pertanyaan & Diskusi",
    )

    return prs


def main() -> None:
    prs = build()
    prs.save(str(OUT))
    print(f"✓ {OUT} berhasil disimpan ({len(prs.slides)} slide)")


if __name__ == "__main__":
    main()
